import zipfile

import numpy as np
import soundfile as sf
from openpyxl import load_workbook

from app.models import Clip, ProcessingSettings, Project
from app.pipeline.export import build_dataset

SENTENCES = [
    ("کیا حال ہے؟ میں ابھی دفتر پہنچا ہوں۔", "Kya haal hai? Main abhi daftar pohancha hoon."),
    ("آپ کہاں جا رہے ہیں؟", "Aap kahan ja rahe hain?"),
    ("مجھے یہ کام آج مکمل کرنا ہے۔", "Mujhe yeh kaam aaj mukammal karna hai."),
]


def make_project(tmp_path, n=3, roman=True, **settings):
    (tmp_path / "clips").mkdir()
    clips = []
    for i in range(n):
        rel = f"clips/c{i}.wav"
        sf.write(str(tmp_path / rel), (0.1 * np.sin(np.arange(22050) / 10)).astype("float32"), 22050, subtype="PCM_16")
        urdu, rom = SENTENCES[i % 3]
        clips.append(Clip(id=f"c{i}", source_id="s", order=i, start_s=0, end_s=1, status="processed",
                          wav=rel, urdu=urdu, roman=rom if roman else "", duration_s=1.0))
    return Project(id="p", name="t", created_at=0, settings=ProcessingSettings(**settings), clips=clips)


def test_numbering_starts_at_start_index(tmp_path):
    report = build_dataset(make_project(tmp_path, start_index=101, end_index=300), tmp_path)
    assert report["ok"] and report["first"] == "urdu_000101.wav" and report["last"] == "urdu_000103.wav"
    names = sorted(p.name for p in (tmp_path / "output/dataset/wavs").iterdir())
    assert names == ["urdu_000101.wav", "urdu_000102.wav", "urdu_000103.wav"]


def test_metadata_csv_is_pipe_separated_utf8_without_bom(tmp_path):
    build_dataset(make_project(tmp_path, start_index=1, end_index=3), tmp_path)
    raw = (tmp_path / "output/dataset/metadata.csv").read_bytes()
    assert not raw.startswith(b"\xef\xbb\xbf")
    lines = raw.decode("utf-8").splitlines()
    assert lines[0] == "urdu_000001.wav|کیا حال ہے؟ میں ابھی دفتر پہنچا ہوں۔|Kya haal hai? Main abhi daftar pohancha hoon."
    assert all(line.count("|") == 2 for line in lines) and len(lines) == 3


def test_xlsx_has_formula_column_and_matches_csv_layout(tmp_path):
    build_dataset(make_project(tmp_path, start_index=1, end_index=3), tmp_path)
    ws = load_workbook(tmp_path / "output/metadata.xlsx").active
    assert [c.value for c in ws[1]] == ["Filename", "Urdu Sentence", "Roman Urdu", "Final"]
    assert ws["D2"].value == '=A2&"|"&B2&"|"&C2' and ws["D4"].value == '=A4&"|"&B4&"|"&C4'
    assert ws["A2"].value == "urdu_000001.wav" and ws["B3"].value == "آپ کہاں جا رہے ہیں؟"


def test_wavs_are_mono_16bit_at_target_rate_and_zip_has_the_structure(tmp_path):
    build_dataset(make_project(tmp_path, start_index=1, end_index=3), tmp_path)
    info = sf.info(str(tmp_path / "output/dataset/wavs/urdu_000002.wav"))
    assert (info.channels, info.samplerate, info.subtype) == (1, 22050, "PCM_16")
    names = set(zipfile.ZipFile(tmp_path / "output/dataset.zip").namelist())
    assert {"dataset/metadata.csv", "dataset/wavs/urdu_000001.wav", "metadata.xlsx"} <= names


def test_missing_roman_blocks_export_unless_forced(tmp_path):
    project = make_project(tmp_path, roman=False, start_index=1, end_index=3)
    blocked = build_dataset(project, tmp_path)
    assert not blocked["written"] and any("Roman" in e for e in blocked["errors"])
    forced = build_dataset(project, tmp_path, force=True)
    assert forced["written"] and not forced["ok"]


def test_clips_beyond_capacity_are_left_out_with_a_warning(tmp_path):
    report = build_dataset(make_project(tmp_path, n=3, start_index=1, end_index=2), tmp_path)
    assert report["count"] == 2 and any("beyond" in w for w in report["warnings"])


def test_without_a_last_file_number_every_kept_clip_is_exported(tmp_path):
    project = make_project(tmp_path, n=5)  # defaults: numbering from 1, no upper limit
    assert project.settings.start_index == 1 and project.settings.end_index is None
    report = build_dataset(project, tmp_path)
    assert report["ok"] and report["count"] == 5 and report["expected"] is None and not report["warnings"]
    assert report["first"] == "urdu_000001.wav" and report["last"] == "urdu_000005.wav"


def test_rejected_clips_are_not_exported_and_numbering_has_no_gaps(tmp_path):
    project = make_project(tmp_path, start_index=1, end_index=3)
    project.clips[1].keep = False
    report = build_dataset(project, tmp_path)
    assert report["count"] == 2 and report["last"] == "urdu_000002.wav"
