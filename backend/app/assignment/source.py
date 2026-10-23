from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import soundfile as sf

from .. import config

NUMBERED = re.compile(r"(\d+)")


@dataclass
class SourceRecording:
    key: str
    wav: Path
    urdu: str
    roman: str


@dataclass
class SourceScan:
    folder: Path
    recordings: list[SourceRecording]
    sample_rates: list[int]
    durations: list[float]
    missing_audio: list[str]
    empty_text: list[str]
    metadata_file: str


class SourceError(Exception):
    pass


def default_folder() -> Path:
    return config.ROOT_DIR / "dataset"


def resolve_folder(text: str) -> Path:
    raw = (text or "").strip().strip('"').strip("'")
    if not raw:
        return default_folder()
    path = Path(raw).expanduser()
    return path if path.is_absolute() else config.ROOT_DIR / path


def natural_key(name: str):
    return [int(part) if part.isdigit() else part.lower() for part in NUMBERED.split(name)]


def find_layout(folder: Path) -> tuple[Path, Path]:
    for base in (folder, folder / "dataset"):
        if (base / "wavs").is_dir():
            return base, base / "wavs"
    raise SourceError(f"No wavs folder found in {folder}. Point this at the folder that holds wavs/ and the metadata file.")


def read_rows(folder: Path, base: Path) -> tuple[list[tuple[str, str, str]], str]:
    csv_path = next((p for p in (base / "metadata.csv", folder / "metadata.csv") if p.is_file()), None)
    if csv_path is not None:
        rows: list[tuple[str, str, str]] = []
        with open(csv_path, encoding="utf-8-sig", newline="") as handle:
            for line in handle.read().splitlines():
                parts = line.split("|")
                if len(parts) >= 3 and parts[0].strip().lower().endswith(".wav"):
                    rows.append((parts[0].strip(), parts[1].strip(), "|".join(parts[2:]).strip()))
        return rows, csv_path.name
    xlsx_path = next((p for p in (folder / "metadata.xlsx", base / "metadata.xlsx", base.parent / "metadata.xlsx") if p.is_file()), None)
    if xlsx_path is None:
        raise SourceError("No metadata.csv or metadata.xlsx found next to the wavs folder.")
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        rows = []
        for row in sheet.iter_rows(min_row=1, values_only=True):
            name = str(row[0] or "").strip()
            if name.lower().endswith(".wav"):
                rows.append((name, str(row[1] or "").strip(), str(row[2] or "").strip()))
    finally:
        workbook.close()
    return rows, xlsx_path.name


def scan(folder: Path) -> SourceScan:
    if not folder.is_dir():
        raise SourceError(f"Folder not found: {folder}")
    base, wavs = find_layout(folder)
    rows, metadata_file = read_rows(folder, base)
    rows.sort(key=lambda row: natural_key(row[0]))
    recordings: list[SourceRecording] = []
    rates: set[int] = set()
    durations: list[float] = []
    missing: list[str] = []
    empty: list[str] = []
    for name, urdu, roman in rows:
        path = wavs / name
        if not path.is_file():
            missing.append(name)
            continue
        info = sf.info(str(path))
        rates.add(info.samplerate)
        durations.append(info.frames / info.samplerate)
        if not urdu.strip() or not roman.strip():
            empty.append(name)
        recordings.append(SourceRecording(key=Path(name).stem, wav=path, urdu=urdu, roman=roman))
    if not recordings:
        raise SourceError("The metadata file has no rows that match a wav in the wavs folder.")
    return SourceScan(folder, recordings, sorted(rates), durations, missing, empty, metadata_file)
