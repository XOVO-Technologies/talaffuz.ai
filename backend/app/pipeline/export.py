"""Write the final dataset: wavs/, metadata.csv (pipe separated, UTF-8), metadata.xlsx, zip."""
from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

import soundfile as sf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from ..models import Project
from .text_ur import clean_roman, normalize_urdu

FORMULA = '=A{r}&"|"&B{r}&"|"&C{r}'


def filename_for(index: int) -> str:
    return f"urdu_{index:06d}.wav"


def select_rows(project: Project) -> tuple[list, list]:
    """Kept clips in recording order, split into (exported, over_capacity)."""
    kept = [c for c in project.clips if c.keep and c.status == "processed" and c.wav]
    kept.sort(key=lambda c: c.order)
    cap = project.settings.capacity
    return (kept, []) if cap is None else (kept[:cap], kept[cap:])


def validate(project: Project, rows: list, extra: list) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    cap = project.settings.capacity
    if not rows:
        errors.append("No kept clips to export.")
    for i, c in enumerate(rows):
        name = filename_for(project.settings.start_index + i)
        if not normalize_urdu(c.urdu, ensure_terminal=False):
            errors.append(f"{name}: the Urdu sentence is empty.")
        if not clean_roman(c.roman):
            errors.append(f"{name}: the Roman Urdu is empty.")
    if rows and cap is not None and len(rows) < cap:
        warnings.append(f"Only {len(rows)} of {cap} clips are ready (urdu_{project.settings.start_index:06d} to "
                        f"urdu_{project.settings.start_index + len(rows) - 1:06d}). Process more clips to reach the target.")
    if extra:
        warnings.append(f"{len(extra)} kept clips are beyond the {cap} slots and are left out of the export.")
    return errors, warnings


def _write_xlsx(path: Path, rows: list[tuple[str, str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "metadata"
    ws.append(["Filename", "Urdu Sentence", "Roman Urdu", "Final"])
    for r, (fn, urdu, roman) in enumerate(rows, start=2):
        ws.append([fn, urdu, roman, FORMULA.format(r=r)])  # Final is a live formula, as the spec asks
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in range(2, len(rows) + 2):
        ws.cell(r, 2).alignment = Alignment(horizontal="right", readingOrder=2)  # right-to-left
        ws.cell(r, 2).font = Font(size=14)
    for col, width in zip("ABCD", (20, 60, 55, 110)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.save(path)


def _zip_dataset(out_dir: Path, zip_path: Path) -> None:
    with zipfile.ZipFile(zip_path, "w") as z:
        for f in sorted((out_dir / "dataset").rglob("*")):
            if f.is_file():
                z.write(f, f.relative_to(out_dir).as_posix(), compress_type=zipfile.ZIP_STORED if f.suffix == ".wav" else zipfile.ZIP_DEFLATED)
        z.write(out_dir / "metadata.xlsx", "metadata.xlsx", compress_type=zipfile.ZIP_DEFLATED)


def check_written(project: Project, out_dir: Path, names: list[str]) -> list[str]:
    """Re-read what was written and verify the format rules from the dataset guide."""
    problems: list[str] = []
    raw = (out_dir / "dataset" / "metadata.csv").read_bytes()
    if raw.startswith(b"\xef\xbb\xbf"):
        problems.append("metadata.csv starts with a UTF-8 BOM.")
    lines = raw.decode("utf-8").splitlines()
    if len(lines) != len(names):
        problems.append(f"metadata.csv has {len(lines)} lines for {len(names)} wav files.")
    for line, name in zip(lines, names):
        parts = line.split("|")
        if len(parts) != 3 or parts[0] != name or not parts[1] or not parts[2]:
            problems.append(f"Bad metadata line for {name}.")
    sr = project.settings.sample_rate
    for name in names:
        info = sf.info(str(out_dir / "dataset" / "wavs" / name))
        if info.channels != 1 or info.samplerate != sr or info.subtype != "PCM_16":
            problems.append(f"{name} is not mono 16-bit PCM at {sr} Hz.")
    return problems


def build_dataset(project: Project, project_dir: Path, force: bool = False) -> dict:
    rows, extra = select_rows(project)
    errors, warnings = validate(project, rows, extra)
    report = {
        "written": False,
        "ok": not errors,
        "errors": errors,
        "warnings": warnings,
        "count": len(rows),
        "expected": project.settings.capacity,  # None when there is no last file number
        "first": filename_for(project.settings.start_index) if rows else None,
        "last": filename_for(project.settings.start_index + len(rows) - 1) if rows else None,
        "total_seconds": round(sum(c.duration_s for c in rows), 1),
        "output_dir": None,
    }
    if (errors and not force) or not rows:
        return report

    out_dir = project_dir / "output"
    if out_dir.exists():
        shutil.rmtree(out_dir)
    wavs = out_dir / "dataset" / "wavs"
    wavs.mkdir(parents=True)

    table: list[tuple[str, str, str]] = []
    for i, c in enumerate(rows):
        name = filename_for(project.settings.start_index + i)
        shutil.copyfile(project_dir / c.wav, wavs / name)
        table.append((name, normalize_urdu(c.urdu, ensure_terminal=False), clean_roman(c.roman)))

    # UTF-8 *without* BOM: training code reads the first filename as-is.
    (out_dir / "dataset" / "metadata.csv").write_bytes(
        ("".join(f"{fn}|{urdu}|{roman}\n" for fn, urdu, roman in table)).encode("utf-8")
    )
    _write_xlsx(out_dir / "metadata.xlsx", table)
    _zip_dataset(out_dir, out_dir / "dataset.zip")

    problems = check_written(project, out_dir, [t[0] for t in table])
    report.update(written=True, ok=not errors and not problems, output_dir=str(out_dir))
    report["errors"] = errors + problems
    return report
